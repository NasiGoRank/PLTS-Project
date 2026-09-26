from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta, timezone
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from api import STORE, app, require_authenticated_user

PAGE_SIZE = 1000
MAX_DAILY_RANGE_DAYS = 3660
MAX_HOURLY_RANGE_DAYS = 366
EXPORT_TIMEZONE = ZoneInfo("Asia/Jakarta")


def _parse_date(value: str, field_name: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=f"{field_name} must use YYYY-MM-DD format") from exc


def _validate_range(start_date: str, end_date: str, resolution: str) -> tuple[date, date]:
    start = _parse_date(start_date, "start_date")
    end = _parse_date(end_date, "end_date")
    if end < start:
        raise HTTPException(status_code=422, detail="end_date must be on or after start_date")

    range_days = (end - start).days + 1
    maximum = MAX_HOURLY_RANGE_DAYS if resolution == "hourly" else MAX_DAILY_RANGE_DAYS
    if range_days > maximum:
        raise HTTPException(
            status_code=422,
            detail=f"{resolution} exports are limited to {maximum} days per download",
        )
    return start, end


def _fetch_pages(build_query) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    while True:
        response = build_query().range(offset, offset + PAGE_SIZE - 1).execute()
        batch = list(response.data or [])
        rows.extend(batch)
        if len(batch) < PAGE_SIZE:
            break
        offset += PAGE_SIZE
    return rows


def _timestamp_parts(
    value: Any,
    target_timezone: timezone | ZoneInfo = EXPORT_TIMEZONE,
) -> tuple[str | None, str | None]:
    """Return separate YYYY-MM-DD and HH:MM:SS values in the requested timezone."""
    if value in (None, ""):
        return None, None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        text = str(value)
        if "T" in text:
            date_part, time_part = text.split("T", 1)
            return date_part or None, time_part or None
        if " " in text:
            date_part, time_part = text.split(" ", 1)
            return date_part or None, time_part or None
        return text, None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    converted = parsed.astimezone(target_timezone)
    return converted.date().isoformat(), converted.time().replace(microsecond=0).isoformat()


def _daily_rows(
    *,
    start: date,
    end: date,
    station_id: str | None,
    platform: str | None,
) -> list[dict[str, Any]]:
    client, config = STORE._require()

    def build_query():
        query = (
            client.table(config.daily_table)
            .select(
                "platform,station_id,station_name,bucket_date,energy_kwh,"
                "revenue_amount,currency,source_scraped_at"
            )
            .gte("bucket_date", start.isoformat())
            .lte("bucket_date", end.isoformat())
            .order("bucket_date")
        )
        if station_id:
            query = query.eq("station_id", station_id)
        if platform:
            query = query.eq("platform", platform)
        return query

    with STORE._lock:
        rows = _fetch_pages(build_query)

    for row in rows:
        scraped_date, scraped_time = _timestamp_parts(row.get("source_scraped_at"))
        row["source_scraped_date_wib"] = scraped_date
        row["source_scraped_time_wib"] = scraped_time
    return rows


def _hourly_utc_bounds(start: date, end: date) -> tuple[datetime, datetime]:
    """Convert inclusive WIB calendar dates into an exclusive UTC query range."""
    start_local = datetime.combine(start, time.min, tzinfo=EXPORT_TIMEZONE)
    end_local = datetime.combine(
        end + timedelta(days=1),
        time.min,
        tzinfo=EXPORT_TIMEZONE,
    )
    return start_local.astimezone(timezone.utc), end_local.astimezone(timezone.utc)


def _hourly_rows(
    *,
    start: date,
    end: date,
    station_id: str | None,
    platform: str | None,
) -> list[dict[str, Any]]:
    client, config = STORE._require()
    start_at, end_exclusive = _hourly_utc_bounds(start, end)

    def build_query():
        query = (
            client.table(config.hourly_table)
            .select(
                "bucket_hour,platform,station_id,station_name,status,current_power_kw,"
                "daily_energy_kwh,monthly_energy_kwh,yearly_energy_kwh,cumulative_energy_kwh,"
                "daily_income,monthly_income,yearly_income,cumulative_income,currency,"
                "station_timezone,source_scraped_at"
            )
            .gte("bucket_hour", start_at.isoformat())
            .lt("bucket_hour", end_exclusive.isoformat())
            .order("bucket_hour")
        )
        if station_id:
            query = query.eq("station_id", station_id)
        if platform:
            query = query.eq("platform", platform)
        return query

    with STORE._lock:
        rows = _fetch_pages(build_query)

    for row in rows:
        hour_date_wib, hour_time_wib = _timestamp_parts(row.get("bucket_hour"))
        reading_date_wib, reading_time_wib = _timestamp_parts(row.get("source_scraped_at"))

        row["bucket_hour_date_wib"] = hour_date_wib
        row["bucket_hour_time_wib"] = hour_time_wib
        row["source_scraped_date_wib"] = reading_date_wib
        row["source_scraped_time_wib"] = reading_time_wib
    return rows


def _csv_response(rows: list[dict[str, Any]], columns: list[tuple[str, str]], filename: str) -> StreamingResponse:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([row.get(key) for key, _ in columns])

    content = "\ufeff" + output.getvalue()
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Cache-Control": "no-store, max-age=0",
        "X-Export-Row-Count": str(len(rows)),
    }
    return StreamingResponse(iter([content]), media_type="text/csv; charset=utf-8", headers=headers)


def _xlsx_response(
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str]],
    filename: str,
) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "History"
    sheet.freeze_panes = "A2"
    sheet.append([label for _, label in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(vertical="center")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for row in rows:
        sheet.append([row.get(key) for key, _ in columns])

    if columns:
        sheet.auto_filter.ref = sheet.dimensions

    sample_rows = rows[:500]
    for index, (key, label) in enumerate(columns, start=1):
        width = len(label)
        for row in sample_rows:
            value = row.get(key)
            if value is not None:
                width = max(width, len(str(value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 34)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Cache-Control": "no-store, max-age=0",
        "X-Export-Row-Count": str(len(rows)),
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _filename_token(value: Any, fallback: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return (text[:80] or fallback).strip("-")


def _export_filename(
    *,
    rows: list[dict[str, Any]],
    start: date,
    end: date,
    resolution: str,
    file_format: str,
    station_id: str | None,
    platform: str | None,
    station_name: str | None,
    generated_at: datetime | None = None,
) -> str:
    if station_id:
        row_station_name = next(
            (str(row.get("station_name")) for row in rows if row.get("station_name")),
            None,
        )
        site_name = station_name or row_station_name or station_id
        platform_token = _filename_token(platform, "Platform") if platform else None
        site_token = _filename_token(site_name, "Selected-Site")
        scope = "Site-" + "-".join(part for part in (platform_token, site_token) if part)
    else:
        scope = "All-Sites"

    exported_at = (generated_at or datetime.now(EXPORT_TIMEZONE)).astimezone(EXPORT_TIMEZONE)
    exported_stamp = exported_at.strftime("%Y-%m-%d_%H-%M-WIB")
    extension = "xlsx" if file_format == "xlsx" else "csv"
    return (
        f"PLTS-History_{scope}_{resolution.title()}_"
        f"{start.isoformat()}_to_{end.isoformat()}_"
        f"Exported-{exported_stamp}.{extension}"
    )


@app.get("/api/history/export")
def export_history(
    start_date: str = Query(
        ...,
        description="Start date in YYYY-MM-DD format; exports use Asia/Jakarta calendar dates",
    ),
    end_date: str = Query(
        ...,
        description="End date in YYYY-MM-DD format; exports use Asia/Jakarta calendar dates",
    ),
    resolution: str = Query(default="hourly", pattern="^(daily|hourly)$"),
    file_format: str = Query(default="csv", alias="format", pattern="^(csv|xlsx)$"),
    station_id: str | None = Query(default=None),
    station_name: str | None = Query(default=None, max_length=160),
    platform: str | None = Query(default=None),
    _: dict[str, Any] = Depends(require_authenticated_user),
) -> StreamingResponse:
    start, end = _validate_range(start_date, end_date, resolution)

    try:
        if resolution == "daily":
            rows = _daily_rows(start=start, end=end, station_id=station_id, platform=platform)
            columns = [
                ("bucket_date", "Date"),
                ("platform", "Platform"),
                ("station_id", "Station ID"),
                ("station_name", "Station Name"),
                ("energy_kwh", "Energy (kWh)"),
                ("revenue_amount", "Revenue"),
                ("currency", "Currency"),
                ("source_scraped_date_wib", "Source Scraped Date (WIB)"),
                ("source_scraped_time_wib", "Source Scraped Time (WIB)"),
            ]
        else:
            rows = _hourly_rows(start=start, end=end, station_id=station_id, platform=platform)
            columns = [
                ("bucket_hour_date_wib", "Hour Date (WIB)"),
                ("bucket_hour_time_wib", "Hour Time (WIB)"),
                ("platform", "Platform"),
                ("station_id", "Station ID"),
                ("station_name", "Station Name"),
                ("status", "Status"),
                ("current_power_kw", "Current Power (kW)"),
                ("daily_energy_kwh", "Daily Energy (kWh)"),
                ("monthly_energy_kwh", "Monthly Energy (kWh)"),
                ("yearly_energy_kwh", "Yearly Energy (kWh)"),
                ("cumulative_energy_kwh", "Cumulative Energy (kWh)"),
                ("daily_income", "Daily Revenue"),
                ("monthly_income", "Monthly Revenue"),
                ("yearly_income", "Yearly Revenue"),
                ("cumulative_income", "Cumulative Revenue"),
                ("currency", "Currency"),
                ("station_timezone", "Station Timezone"),
                ("source_scraped_date_wib", "Latest Reading Date (WIB)"),
                ("source_scraped_time_wib", "Latest Reading Time (WIB)"),
            ]
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail=f"History export is unavailable: {type(exc).__name__}: {exc}",
        ) from exc

    filename = _export_filename(
        rows=rows,
        start=start,
        end=end,
        resolution=resolution,
        file_format=file_format,
        station_id=station_id,
        platform=platform,
        station_name=station_name,
    )
    if file_format == "xlsx":
        return _xlsx_response(rows, columns, filename)
    return _csv_response(rows, columns, filename)
